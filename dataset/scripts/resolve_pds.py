#!/usr/bin/env python3
"""
resolve_pds.py — Resuelve DID y PDS de cada cuenta de Bluesky.

El handle no dice dónde vive la cuenta: un dominio propio (aemet.es) puede estar
alojado en Bluesky, en Eurosky o en un PDS propio. Lo que lo determina es el DID
y el endpoint `#atproto_pds` de su documento DID. Este script lo resuelve y lo
guarda en datosfinales.xlsx (columnas "Bluesky DID" y "Bluesky PDS"); export.py
marca como Eurosky las cuentas cuyo PDS es eurosky.social.

Uso:
    python3 dataset/scripts/resolve_pds.py            # todas las cuentas
    python3 dataset/scripts/resolve_pds.py --only-missing
    python3 dataset/scripts/resolve_pds.py --dry-run
"""
import sys, os, json, time, argparse, urllib.request, urllib.parse
from urllib.parse import urlparse
import openpyxl

SCRIPT_DIR  = os.path.dirname(os.path.abspath(__file__))
DF_FILE     = os.path.join(os.path.dirname(SCRIPT_DIR), "datosfinales.xlsx")
C_NOMBRE=2; C_BS=5; C_BS_DID=23; C_BS_PDS=24
RESOLVE = "https://public.api.bsky.app/xrpc/com.atproto.identity.resolveHandle?handle="

def _get(url):
    req = urllib.request.Request(url, headers={"User-Agent": "saldeahi-dataset/1.0"})
    with urllib.request.urlopen(req, timeout=10) as r:
        return json.loads(r.read())

def resolve_account(handle):
    """(did, pds_host) para un handle; (None, None) si no se puede resolver.
    Lanza OSError ante fallos de red para no confundirlos con 'no existe'."""
    h = handle.lstrip("@").lower().strip()
    if "." not in h: h += ".bsky.social"
    try:
        did = _get(RESOLVE + urllib.parse.quote(h)).get("did")
    except urllib.error.HTTPError as e:
        if e.code in (400, 404): return None, None
        raise
    if not did: return None, None
    if did.startswith("did:plc:"):
        doc = _get("https://plc.directory/" + did)
    elif did.startswith("did:web:"):
        doc = _get(f"https://{did[8:]}/.well-known/did.json")
    else:
        return did, None
    for svc in doc.get("service", []):
        if svc.get("id", "").endswith("#atproto_pds") or svc.get("type") == "AtprotoPersonalDataServer":
            return did, urlparse(svc["serviceEndpoint"]).hostname
    return did, None

def is_eurosky_pds(pds):
    return bool(pds) and (pds == "eurosky.social" or pds.endswith(".eurosky.social"))

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only-missing", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    wb = openpyxl.load_workbook(DF_FILE); ws = wb.active
    ws.cell(1, C_BS_DID).value = "Bluesky DID"; ws.cell(1, C_BS_PDS).value = "Bluesky PDS"
    cache, n_ok, n_fail = {}, 0, 0
    for row in range(2, ws.max_row + 1):
        handle = ws.cell(row, C_BS).value
        if not handle: 
            ws.cell(row, C_BS_DID).value = None; ws.cell(row, C_BS_PDS).value = None
            continue
        if args.only_missing and ws.cell(row, C_BS_PDS).value: continue
        key = handle.lower().strip()
        if key not in cache:
            try: cache[key] = resolve_account(handle)
            except Exception as e:
                print(f"  ! {handle}: {e}"); n_fail += 1; time.sleep(0.5); continue
            time.sleep(0.1)
        did, pds = cache[key]
        if did is None:
            # Cuenta inexistente o handle caído: se conserva lo anterior, si lo había.
            print(f"  ? {handle}: no resuelve"); n_fail += 1; continue
        ws.cell(row, C_BS_DID).value = did; ws.cell(row, C_BS_PDS).value = pds
        n_ok += 1
        if is_eurosky_pds(pds): print(f"  ★ Eurosky: {ws.cell(row, C_NOMBRE).value} ({handle})")
    print(f"\nResueltas: {n_ok} · sin resolver: {n_fail}")
    if args.dry_run: print("DRY-RUN: no se guarda."); return
    wb.save(DF_FILE); print("datosfinales.xlsx actualizado.")

if __name__ == "__main__":
    main()
